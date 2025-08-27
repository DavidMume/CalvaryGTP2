import base64
import io
import json
import logging
import os
import uuid
from typing import List

import azure.functions as func
from pptx import Presentation
from pptx.util import Pt

from .models import PresentationRequest, Slide

logger = logging.getLogger("generate_pptx")

app = func.FunctionApp(http_auth_level=func.AuthLevel.ANONYMOUS)


def _load_base_presentation() -> Presentation:
    """Load a base presentation from BRAND_TEMPLATE_PATH if available."""
    template_path = os.getenv("BRAND_TEMPLATE_PATH")
    if template_path and os.path.exists(template_path):
        try:
            return Presentation(template_path)
        except Exception:  # pragma: no cover - defensive
            logger.warning("Failed to load template", extra={"template": template_path})
    return Presentation()


def build_pptx(slides: List[Slide], style: str) -> bytes:
    """Build a PowerPoint presentation."""
    prs = _load_base_presentation()
    layout = prs.slide_layouts[1]

    style_fonts = {
        "minimalist": (Pt(32), Pt(18)),
        "corporate": (Pt(36), Pt(20)),
        "colorful": (Pt(40), Pt(24)),
    }
    title_size, bullet_size = style_fonts.get(style, style_fonts["minimalist"])

    for slide in slides:
        sld = prs.slides.add_slide(layout)
        title_placeholder = sld.shapes.title
        content = sld.placeholders[1]

        title_placeholder.text = slide.title
        title_placeholder.text_frame.paragraphs[0].font.size = title_size

        tf = content.text_frame
        tf.clear()
        for bullet in slide.bullets:
            p = tf.add_paragraph()
            p.text = bullet
            p.font.size = bullet_size
            p.level = 0

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@app.function_name(name="generate_pptx")
@app.route(route="generate-pptx", methods=["POST"])
def generate_pptx(req: func.HttpRequest) -> func.HttpResponse:
    """HTTP trigger to generate a PPTX file based on JSON payload."""
    request_id = req.headers.get("x-request-id", str(uuid.uuid4()))

    body = req.get_body()
    if len(body) > 1_000_000:  # 1 MB limit
        logger.warning("Payload too large", extra={"request_id": request_id})
        return func.HttpResponse(
            json.dumps({"error": "Payload too large"}),
            status_code=413,
            mimetype="application/json",
        )

    try:
        payload = PresentationRequest.model_validate_json(body)
    except Exception as exc:  # pydantic ValidationError
        logger.warning("Validation error", extra={"request_id": request_id})
        return func.HttpResponse(
            json.dumps({"error": "Invalid payload", "details": getattr(exc, "errors", lambda: [])()}),
            status_code=400,
            mimetype="application/json",
        )

    pptx_bytes = build_pptx(payload.slides, payload.style)
    pptx_b64 = base64.b64encode(pptx_bytes).decode("utf-8")

    logger.info(
        "Generated PPTX",
        extra={"request_id": request_id, "slideCount": len(payload.slides)},
    )

    response = {
        "status": "ok",
        "fileName": payload.fileName,
        "slideCount": len(payload.slides),
        "pptxBase64": pptx_b64,
    }
    return func.HttpResponse(
        json.dumps(response), status_code=200, mimetype="application/json"
    )
